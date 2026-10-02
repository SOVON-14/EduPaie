import contextlib
import io
import os
import sys
import tempfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

INK = "173B36"
ACCENT = "287A5D"
MUTED = "64736F"
PALE = "EDF4F0"
BORDER = "D6E2DC"
WHITE = "FFFFFF"


def capture_screenshots(output_dir):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

    import database.database as database
    from PySide6.QtWidgets import QApplication

    from data.seed import seed
    from ui.main_window import MainWindow
    from ui.student_detail import StudentDetailDialog

    previous_db_path = database.DB_PATH
    captures = {
        "dashboard": output_dir / "dashboard.png",
        "students": output_dir / "students.png",
        "history": output_dir / "history.png",
    }

    try:
        with tempfile.TemporaryDirectory(prefix="edupaie_doc_data_") as temp_dir:
            database.DB_PATH = Path(temp_dir) / "documentation.db"
            with contextlib.redirect_stdout(io.StringIO()):
                seed()

            app = QApplication.instance() or QApplication(["EduPaie documentation"])
            app.setStyle("Fusion")
            window = MainWindow()
            window.resize(1180, 700)
            window.show()
            app.processEvents()
            if not window.grab().save(str(captures["dashboard"])):
                raise RuntimeError("Impossible de capturer le tableau de bord.")

            window.tabs.setCurrentWidget(window.students_tab)
            app.processEvents()
            if not window.grab().save(str(captures["students"])):
                raise RuntimeError("Impossible de capturer la liste des élèves.")

            student = next(
                student
                for student in window.student_service.get_all_students()
                if student["total_paye"] > 0
            )
            detail = StudentDetailDialog(
                student["id"],
                window.student_service,
                window.payment_service,
                window.receipt_service,
            )
            detail.resize(960, 620)
            detail.show()
            app.processEvents()
            if not detail.grab().save(str(captures["history"])):
                raise RuntimeError("Impossible de capturer l'historique des paiements.")

            detail.close()
            window.close()
            app.processEvents()
    finally:
        database.DB_PATH = previous_db_path

    return captures


def _set_docx_cell_fill(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _format_docx_table(table, header=True):
    from docx.shared import Pt, RGBColor

    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = 1
            if header and row_index == 0:
                _set_docx_cell_fill(cell, INK)
            elif row_index % 2 == 0:
                _set_docx_cell_fill(cell, PALE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(8.5)
                    if header and row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string(WHITE)


def _add_docx_footer(section):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("EduPaie | Documentation technique | ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def create_docx(captures, path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.75)
    section.footer_distance = Cm(0.75)
    _add_docx_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for style_name, size, color in (("Heading 1", 16, INK), ("Heading 2", 11, ACCENT)):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = "EDUPAIE  /  DOCUMENTATION TECHNIQUE"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(ACCENT)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(30)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EDUPAIE\nArchitecture et choix techniques")
    run.font.name = "Calibri"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)
    subtitle = doc.add_paragraph("Documentation courte | Version 0.1.0 | 2 octobre 2026")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(MUTED)

    doc.add_heading("1. Présentation et architecture", level=1)
    doc.add_paragraph(
        "EduPaie est une application de bureau destinée au suivi des frais de scolarité. "
        "Elle gère les élèves, les paiements, les soldes et la production de reçus PDF. "
        "L’application est structurée en couches afin de séparer l’interface, les règles métier et l’accès aux données."
    )
    architecture = doc.add_table(rows=1, cols=2)
    architecture.rows[0].cells[0].text = "Couche"
    architecture.rows[0].cells[1].text = "Responsabilité"
    architecture_rows = [
        ("Interface — PySide6", "Fenêtre principale, tableaux, formulaires et dialogues."),
        ("Services", "Règles métier : élèves, paiements, statistiques et reçus."),
        ("Repositories", "Requêtes SQLite pour lire et modifier les élèves et paiements."),
        ("Base — SQLite", "Persistance locale, contraintes, index et migrations de schéma."),
        ("Sortie PDF — ReportLab", "Mise en page des reçus de paiement générés par ReceiptService."),
    ]
    for layer, purpose in architecture_rows:
        cells = architecture.add_row().cells
        cells[0].text = layer
        cells[1].text = purpose
    _format_docx_table(architecture)
    doc.add_paragraph(
        "Flux courant : l’interface appelle un service; le service applique les validations et délègue "
        "la persistance au repository. ReceiptService assemble les informations et produit le PDF."
    )

    doc.add_page_break()
    doc.add_heading("2. Choix techniques et données", level=1)
    choices = doc.add_table(rows=1, cols=3)
    for index, text in enumerate(("Technologie", "Usage", "Justification observée dans le projet")):
        choices.rows[0].cells[index].text = text
    for row in [
        ("Python", "Logique applicative", "Écosystème simple à maintenir et partagé par les couches métier."),
        ("PySide6", "Interface de bureau", "Fournit les fenêtres natives, tableaux, onglets et dialogues."),
        ("SQLite", "Stockage local", "Base fichier, sans serveur à administrer; transactions et contraintes SQL."),
        ("ReportLab", "Reçus PDF", "Génération de documents paginés avec texte et tableaux."),
        ("PyInstaller + Inno Setup", "Distribution Windows", "Embarque Python puis fournit installation, raccourcis et désinstallation."),
    ]:
        cells = choices.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = text
    _format_docx_table(choices)

    doc.add_heading("Modèle de données", level=2)
    model = doc.add_table(rows=1, cols=3)
    for index, text in enumerate(("Table", "Rôle", "Relation / contraintes principales")):
        model.rows[0].cells[index].text = text
    for row in [
        ("students", "Élèves et montant annuel dû.", "Unicité nom/prénom/classe/année; montant strictement positif."),
        ("payments", "Historique des versements.", "Chaque paiement référence un élève; suppression en cascade."),
    ]:
        cells = model.add_row().cells
        for index, text in enumerate(row):
            cells[index].text = text
    _format_docx_table(model)
    doc.add_paragraph(
        "Relation : un élève peut avoir zéro à plusieurs paiements; chaque paiement appartient à un seul élève. "
        "Les montants sont des entiers en FCFA, le numéro de reçu est unique et le mode de paiement est contrôlé par SQLite."
    )
    doc.add_paragraph(
        "Fiabilité : les connexions activent les clés étrangères et sont fermées systématiquement; "
        "le schéma est versionné via PRAGMA user_version et les mises à niveau sont gérées dans database.py."
    )

    doc.add_page_break()
    doc.add_heading("3. Écrans principaux", level=1)
    doc.add_paragraph(
        "Les captures ci-dessous proviennent de l’application exécutée hors écran avec un jeu de démonstration temporaire. "
        "Aucune donnée de la base locale réelle n’est incluse."
    )
    for key, caption in (
        ("dashboard", "Figure 1 — Tableau de bord : indicateurs globaux et synthèse par classe."),
        ("students", "Figure 2 — Élèves : recherche, filtres de classe et statut, actions sur la sélection."),
    ):
        doc.add_picture(str(captures[key]), width=Inches(6.35))
        paragraph = doc.add_paragraph(caption)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.runs[0].italic = True
        paragraph.runs[0].font.size = Pt(8)

    doc.add_page_break()
    doc.add_heading("4. Historique et limites connues", level=1)
    doc.add_picture(str(captures["history"]), width=Inches(5.65))
    caption = doc.add_paragraph("Figure 3 — Détail d’un élève et historique des paiements.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(8)

    doc.add_heading("Limites actuelles", level=2)
    for item in (
        "Application de bureau Windows; l’installateur configuré cible Windows x64.",
        "SQLite est local au profil Windows (%LOCALAPPDATA%\\EduPaie); la synchronisation réseau entre postes n’est pas implémentée.",
        "Le flux actuel ne présente pas d’écran de connexion ni de rôles utilisateurs.",
        "Aucune sauvegarde automatique depuis l’interface n’est prévue; les sauvegardes doivent être organisées séparément.",
        "Les reçus contiennent l’identité EduPaie, mais aucun nom d’établissement, logo ou coordonnées personnalisés configurables.",
    ):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph(
        "Les captures et le jeu de données utilisés pour ce document sont synthétiques. "
        "Le DOCX est la source entièrement éditable; le PDF conserve du texte sélectionnable et des images intégrées."
    )

    doc.core_properties.title = "EduPaie - Architecture et choix techniques"
    doc.core_properties.subject = "Documentation technique courte"
    doc.core_properties.author = "EduPaie"
    doc.save(path)


def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.enums import TA_CENTER

    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("DocTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=23, leading=27, textColor=colors.HexColor("#" + INK), alignment=TA_CENTER, spaceAfter=8),
        "subtitle": ParagraphStyle("DocSubtitle", parent=base["Normal"], fontName="Helvetica", fontSize=9, leading=13, textColor=colors.HexColor("#" + MUTED), alignment=TA_CENTER, spaceAfter=14),
        "h1": ParagraphStyle("DocH1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#" + INK), spaceBefore=4, spaceAfter=7),
        "h2": ParagraphStyle("DocH2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#" + ACCENT), spaceBefore=5, spaceAfter=4),
        "body": ParagraphStyle("DocBody", parent=base["BodyText"], fontName="Helvetica", fontSize=8.7, leading=11.5, textColor=colors.HexColor("#293833"), spaceAfter=5),
        "small": ParagraphStyle("DocSmall", parent=base["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9, textColor=colors.HexColor("#" + MUTED), spaceAfter=3),
        "table_head": ParagraphStyle("DocTableHead", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=7.7, leading=9.4, textColor=colors.white),
        "table": ParagraphStyle("DocTable", parent=base["BodyText"], fontName="Helvetica", fontSize=7.5, leading=9.3, textColor=colors.HexColor("#293833")),
        "caption": ParagraphStyle("DocCaption", parent=base["BodyText"], fontName="Helvetica-Oblique", fontSize=7.5, leading=9, textColor=colors.HexColor("#" + MUTED), alignment=TA_CENTER, spaceAfter=5),
        "bullet": ParagraphStyle("DocBullet", parent=base["BodyText"], fontName="Helvetica", fontSize=8, leading=10, leftIndent=11, firstLineIndent=-7, textColor=colors.HexColor("#293833"), spaceAfter=4),
    }


def _pdf_table(rows, widths, header=True):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#" + BORDER)),
        ("INNERGRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#" + BORDER)),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    if header:
        commands.extend([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + INK)),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + PALE)]),
        ])
    table.setStyle(TableStyle(commands))
    return table


def create_pdf(captures, path):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer

    styles = _pdf_styles()
    story = []
    story.extend([
        Spacer(1, 0.75*cm),
        Paragraph("EDUPAIE", styles["title"]),
        Paragraph("Architecture et choix techniques", styles["title"]),
        Paragraph("Documentation courte | Version 0.1.0 | 2 octobre 2026", styles["subtitle"]),
        Paragraph("1. Présentation et architecture", styles["h1"]),
        Paragraph(
            "EduPaie est une application de bureau destinée au suivi des frais de scolarité. "
            "Elle gère les élèves, les paiements, les soldes et les reçus PDF. Son architecture sépare "
            "l’interface, les règles métier, l’accès aux données et la génération documentaire.",
            styles["body"],
        ),
    ])
    arch_rows = [[Paragraph("Couche", styles["table_head"]), Paragraph("Rôle dans le projet", styles["table_head"])]]
    for layer, purpose in [
        ("Interface — PySide6", "Fenêtre principale, tableaux, formulaires et dialogues."),
        ("Services", "Règles métier pour élèves, paiements, statistiques et reçus."),
        ("Repositories", "Requêtes SQLite pour lire et modifier les données."),
        ("SQLite", "Persistance locale, contraintes, index et migrations versionnées."),
        ("ReportLab", "Mise en page et génération des reçus PDF."),
    ]:
        arch_rows.append([Paragraph(layer, styles["table"]), Paragraph(purpose, styles["table"])])
    story.extend([
        _pdf_table(arch_rows, [4.1*cm, 13.5*cm]),
        Spacer(1, 0.25*cm),
        Paragraph(
            "Flux principal : interface → service → repository → SQLite. ReceiptService rassemble les données "
            "du paiement et utilise ReportLab pour produire le justificatif.", styles["body"]
        ),
        Paragraph("2. Choix techniques et données", styles["h1"]),
    ])

    choices = [[Paragraph(label, styles["table_head"]) for label in ("Technologie", "Usage", "Justification")]]
    for row in [
        ("Python", "Logique", "Même langage dans les couches applicatives; écosystème courant."),
        ("PySide6", "Interface", "Fenêtres de bureau, tableaux, onglets et formulaires."),
        ("SQLite", "Données", "Base fichier locale, sans serveur; transactions et contraintes SQL."),
        ("ReportLab", "Reçus", "Création de PDF structurés et imprimables."),
        ("PyInstaller / Inno Setup", "Distribution", "Application Windows autonome et installation avec raccourcis."),
    ]:
        choices.append([Paragraph(cell, styles["table"]) for cell in row])
    story.extend([
        _pdf_table(choices, [3.3*cm, 3.0*cm, 11.3*cm]),
        Spacer(1, 0.2*cm),
        Paragraph("Modèle de données", styles["h2"]),
    ])
    model = [[Paragraph(label, styles["table_head"]) for label in ("Table", "Données et contraintes")]]
    model.extend([
        [Paragraph("students", styles["table"]), Paragraph("Élèves; unicité nom/prénom/classe/année; montant total positif.", styles["table"])],
        [Paragraph("payments", styles["table"]), Paragraph("Versements; chaque paiement référence un élève; numéro de reçu unique; montant positif.", styles["table"])],
    ])
    story.extend([
        _pdf_table(model, [3.2*cm, 14.4*cm]),
        Spacer(1, 0.15*cm),
        Paragraph(
            "Relation 1 à N : un élève peut avoir zéro ou plusieurs paiements; un paiement appartient à un seul élève. "
            "Les modes sont contrôlés en base et la suppression d’un élève supprime ses paiements associés.", styles["body"]
        ),
        Paragraph(
            "Les connexions SQLite activent les clés étrangères et sont toujours fermées. La version du schéma est "
            "suivie par PRAGMA user_version; database.py contient les migrations.", styles["body"]
        ),
        PageBreak(),
        Paragraph("3. Écrans principaux", styles["h1"]),
        Paragraph(
            "Captures réalisées depuis l’application, avec une base temporaire de démonstration; aucune donnée réelle n’a été utilisée.",
            styles["small"],
        ),
    ])

    for key, caption in (
        ("dashboard", "Figure 1 — Tableau de bord : indicateurs globaux et statistiques par classe."),
        ("students", "Figure 2 — Élèves : recherche, filtres par classe/statut et actions sur la sélection."),
    ):
        image = Image(str(captures[key]), width=17.4*cm)
        image._restrictSize(17.4*cm, 10.1*cm)
        story.extend([image, Paragraph(caption, styles["caption"]), Spacer(1, 0.1*cm)])

    story.extend([
        PageBreak(),
        Paragraph("4. Historique et limites connues", styles["h1"]),
    ])
    history_image = Image(str(captures["history"]), width=14.1*cm)
    history_image._restrictSize(14.1*cm, 9.3*cm)
    story.extend([
        history_image,
        Paragraph("Figure 3 — Détail d’un élève et historique de ses paiements.", styles["caption"]),
        Paragraph("Limites actuelles", styles["h2"]),
    ])
    limits = [
        "Application de bureau Windows; l’installateur cible Windows x64.",
        "SQLite est local au profil Windows (%LOCALAPPDATA%/EduPaie); aucune synchronisation réseau entre postes n’est implémentée.",
        "Le flux actuel ne comporte pas d’écran de connexion ni de rôles utilisateurs.",
        "La sauvegarde automatique depuis l’interface n’est pas prévue; les sauvegardes sont à organiser séparément.",
        "Les reçus portent l’identité EduPaie, sans coordonnées ni identité d’établissement configurables.",
    ]
    story.extend(Paragraph("• " + item, styles["bullet"]) for item in limits)
    story.append(Spacer(1, 0.1*cm))
    story.append(Paragraph(
        "Les captures utilisent des données synthétiques. Le PDF contient du texte sélectionnable; "
        "le DOCX fourni avec lui est la version la plus simple à modifier.", styles["small"]
    ))

    def draw_page(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#" + BORDER))
        canvas.setLineWidth(0.5)
        canvas.line(document.leftMargin, A4[1] - 1.25*cm, A4[0] - document.rightMargin, A4[1] - 1.25*cm)
        canvas.setFont("Helvetica-Bold", 7.5)
        canvas.setFillColor(colors.HexColor("#" + ACCENT))
        canvas.drawString(document.leftMargin, A4[1] - 0.95*cm, "EDUPAIE  /  DOCUMENTATION TECHNIQUE")
        canvas.line(document.leftMargin, 1.15*cm, A4[0] - document.rightMargin, 1.15*cm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(colors.HexColor("#" + MUTED))
        canvas.drawString(document.leftMargin, 0.82*cm, "Version 0.1.0 | 2 octobre 2026")
        canvas.drawRightString(A4[0] - document.rightMargin, 0.82*cm, f"Page {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=1.7*cm, leftMargin=1.7*cm,
        topMargin=1.65*cm, bottomMargin=1.5*cm,
        title="EduPaie - Architecture et choix techniques",
        author="EduPaie",
    )
    document.build(story, onFirstPage=draw_page, onLaterPages=draw_page)


def main():
    output_dir = ROOT
    docx_path = output_dir / "Documentation_Technique_EduPaie.docx"
    pdf_path = output_dir / "Documentation_Technique_EduPaie.pdf"

    with tempfile.TemporaryDirectory(prefix="edupaie_doc_captures_") as temp_dir:
        captures = capture_screenshots(Path(temp_dir))
        create_docx(captures, docx_path)
        create_pdf(captures, pdf_path)

    print(f"DOCX editable : {docx_path}")
    print(f"PDF : {pdf_path}")


if __name__ == "__main__":
    main()
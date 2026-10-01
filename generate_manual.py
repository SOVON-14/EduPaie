from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_user_manual():
    """Crée un manuel utilisateur au format Word"""
    doc = Document()
    
    # Titre principal
    title = doc.add_heading('MANUEL UTILISATEUR - EDU PAIE', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Sous-titre
    subtitle = doc.add_paragraph('Système de Gestion des Frais de Scolarité')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    
    doc.add_paragraph()  # Espace
    
    # Introduction
    doc.add_heading('1. INTRODUCTION', level=1)
    doc.add_paragraph(
        'EduPaie est une application de gestion des frais de scolarité pour les établissements '
        'scolaires togolais (Collège et Lycée). Elle permet de gérer les élèves, '
        'enregistrer les paiements, calculer automatiquement les soldes et générer des reçus.'
    )
    
    # Section 2 : Enregistrer un élève
    doc.add_heading('2. ENREGISTRER UN ÉLÈVE', level=1)
    
    doc.add_heading('2.1. Accéder au formulaire', level=2)
    doc.add_paragraph(
        'Pour ajouter un nouvel élève, suivez ces étapes :',
        style='List Number'
    )
    doc.add_paragraph(
        'Lancez l\'application EduPaie en double-cliquant sur le fichier main.py',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Cliquez sur le bouton "Ajouter un élève" en bas de la fenêtre principale',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Une fenêtre de dialogue s\'ouvrira',
        style='List Bullet'
    )
    
    doc.add_heading('2.2. Remplir les informations', level=2)
    doc.add_paragraph('Remplissez les champs suivants :')
    
    table = doc.add_table(rows=6, cols=2)
    table.style = 'Table Grid'
    
    # En-têtes
    table.rows[0].cells[0].text = 'Champ'
    table.rows[0].cells[1].text = 'Description'
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    # Contenu
    table.rows[1].cells[0].text = 'Nom *'
    table.rows[1].cells[1].text = 'Nom de famille de l\'élève (ex: KOFFI)'
    
    table.rows[2].cells[0].text = 'Prénom *'
    table.rows[2].cells[1].text = 'Prénom de l\'élève (ex: Yao)'
    
    table.rows[3].cells[0].text = 'Classe *'
    table.rows[3].cells[1].text = 'Sélectionnez ou saisissez la classe (ex: 6ème A, 2nde B, Terminale C)'
    
    table.rows[4].cells[0].text = 'Année scolaire *'
    table.rows[4].cells[1].text = 'Format YYYY-YYYY (ex: 2024-2025)'
    
    table.rows[5].cells[0].text = 'Montant total *'
    table.rows[5].cells[1].text = 'Montant des frais de scolarité en FCFA'
    
    doc.add_paragraph()
    doc.add_paragraph('* = Champs obligatoires')
    
    doc.add_heading('2.3. Valider et enregistrer', level=2)
    doc.add_paragraph(
        'Cliquez sur le bouton "OK" pour enregistrer l\'élève. '
        'Si un champ est invalide ou manquant, un message d\'erreur s\'affichera.'
    )
    
    # Section 3 : Enregistrer un paiement
    doc.add_heading('3. ENREGISTRER UN PAIEMENT', level=1)
    
    doc.add_heading('3.1. Sélectionner l\'élève', level=2)
    doc.add_paragraph(
        'Pour enregistrer un paiement, vous devez d\'abord sélectionner un élève :',
        style='List Number'
    )
    doc.add_paragraph(
        'Allez dans l\'onglet "Élèves"',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Cliquez sur l\'élève souhaité dans la liste',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Cliquez sur le bouton "Enregistrer paiement"',
        style='List Bullet'
    )
    
    doc.add_heading('3.2. Remplir les informations de paiement', level=2)
    doc.add_paragraph('La fenêtre affiche :')
    doc.add_paragraph('• Les informations de l\'élève sélectionné')
    doc.add_paragraph('• Le solde restant à payer')
    doc.add_paragraph()
    doc.add_paragraph('Remplissez les champs suivants :')
    
    table2 = doc.add_table(rows=2, cols=2)
    table2.style = 'Table Grid'
    
    table2.rows[0].cells[0].text = 'Champ'
    table2.rows[0].cells[1].text = 'Description'
    for cell in table2.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    
    table2.rows[1].cells[0].text = 'Montant *'
    table2.rows[1].cells[1].text = 'Montant du paiement en FCFA (ne peut pas dépasser le solde restant)'
    
    doc.add_paragraph()
    doc.add_paragraph('Sélectionnez le mode de paiement :')
    doc.add_paragraph('• Espèces')
    doc.add_paragraph('• Chèque')
    doc.add_paragraph('• Virement')
    doc.add_paragraph('• Mobile Money')
    
    doc.add_heading('3.3. Validation automatique', level=2)
    doc.add_paragraph(
        'Le système vérifie automatiquement que le montant ne dépasse pas le solde restant. '
        'Si le montant est trop élevé, un avertissement s\'affiche et le bouton "Enregistrer" sera désactivé.'
    )
    
    doc.add_heading('3.4. Générer le reçu', level=2)
    doc.add_paragraph(
        'Après l\'enregistrement du paiement, une boîte de dialogue vous demandera '
        'si vous souhaitez générer le reçu PDF immédiatement.'
    )
    
    # Section 4 : Imprimer un reçu
    doc.add_heading('4. IMPRIMER UN REÇU', level=1)
    
    doc.add_heading('4.1. Depuis l\'enregistrement du paiement', level=2)
    doc.add_paragraph(
        'Lors de l\'enregistrement d\'un paiement, choisissez "Oui" '
        'lorsqu\'on vous demande si vous souhaitez générer le reçu PDF.'
    )
    
    doc.add_heading('4.2. Depuis l\'historique des paiements', level=2)
    doc.add_paragraph(
        'Pour imprimer un reçu existant :',
        style='List Number'
    )
    doc.add_paragraph(
        'Double-cliquez sur un élève dans la liste ou cliquez sur "Voir détails"',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Dans l\'onglet "Historique des paiements", sélectionnez le paiement souhaité',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Cliquez sur le bouton "Imprimer le reçu"',
        style='List Bullet'
    )
    doc.add_paragraph(
        'Le PDF sera généré et vous pourrez choisir de l\'ouvrir pour l\'imprimer',
        style='List Bullet'
    )
    
    doc.add_heading('4.3. Contenu du reçu', level=2)
    doc.add_paragraph('Chaque reçu contient :')
    doc.add_paragraph('• Numéro de reçu unique')
    doc.add_paragraph('• Date du paiement')
    doc.add_paragraph('• Montant payé')
    doc.add_paragraph('• Mode de paiement')
    doc.add_paragraph('• Informations de l\'élève (nom, prénom, classe)')
    doc.add_paragraph('• Solde restant après paiement')
    doc.add_paragraph('• Statut (Soldé / Partiellement payé / Non payé)')
    
    # Section 5 : Tableau de bord
    doc.add_heading('5. TABLEAU DE BORD', level=1)
    doc.add_paragraph(
        'L\'onglet "Tableau de bord" affiche les statistiques globales :'
    )
    doc.add_paragraph('• Nombre total d\'élèves')
    doc.add_paragraph('• Total encaissé')
    doc.add_paragraph('• Total restant dû')
    doc.add_paragraph('• Nombre d\'élèves soldés')
    doc.add_paragraph('• Nombre d\'élèves partiellement payés')
    doc.add_paragraph('• Nombre d\'élèves non payés')
    doc.add_paragraph('• Statistiques par classe')
    
    # Section 6 : Classes disponibles
    doc.add_heading('6. CLASSES DISPONIBLES', level=1)
    doc.add_paragraph('L\'application gère les classes suivantes :')
    doc.add_paragraph('COLLÈGE :')
    doc.add_paragraph('• 6ème A, 6ème B, 6ème C')
    doc.add_paragraph('• 5ème A, 5ème B, 5ème C')
    doc.add_paragraph('• 4ème A, 4ème B, 4ème C')
    doc.add_paragraph('• 3ème A, 3ème B, 3ème C')
    doc.add_paragraph('LYCÉE :')
    doc.add_paragraph('• 2nde A, 2nde B, 2nde C')
    doc.add_paragraph('• 1ère A, 1ère B, 1ère C')
    doc.add_paragraph('• Terminale A, Terminale B, Terminale C')
    doc.add_paragraph()
    doc.add_paragraph('Vous pouvez également saisir une classe personnalisée si nécessaire.')
    
    # Section 7 : Conseils
    doc.add_heading('7. CONSEILS ET BONNES PRATIQUES', level=1)
    doc.add_paragraph('• Vérifiez toujours le solde restant avant d\'enregistrer un paiement')
    doc.add_paragraph('• Utilisez la fonction de recherche pour trouver rapidement un élève')
    doc.add_paragraph('• Utilisez les filtres par classe et statut pour organiser votre liste')
    doc.add_paragraph('• Consultez régulièrement le tableau de bord pour suivre les encaissements')
    doc.add_paragraph('• Conservez une copie des reçus PDF pour vos archives')
    
    # Footer
    doc.add_paragraph()
    doc.add_paragraph()
    footer = doc.add_paragraph('© 2024 - EduPaie - Système de Gestion des Frais de Scolarité')
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].italic = True
    
    # Sauvegarder le document
    doc.save('Manuel_Utilisateur_EduPaie.docx')
    print('Manuel utilisateur créé avec succès : Manuel_Utilisateur_EduPaie.docx')

if __name__ == '__main__':
    create_user_manual()

from openpyxl import load_workbook
from repositories.student_repository import StudentRepository
from typing import List, Dict, Any, Tuple
import logging

logger = logging.getLogger(__name__)


class ImportService:
    """Service pour l'import de données depuis des fichiers Excel"""
    
    def __init__(self):
        self.student_repository = StudentRepository()
    
    def import_students_from_excel(self, file_path: str) -> Tuple[int, int, List[str]]:
        """Importe des élèves depuis un fichier Excel.
        
        Args:
            file_path: Chemin vers le fichier Excel
            
        Returns:
            Tuple (nombre d'élèves importés, nombre d'erreurs, liste des messages d'erreur)
            
        Raises:
            FileNotFoundError: Si le fichier n'existe pas
            ValueError: Si le format du fichier est invalide
        """
        try:
            workbook = load_workbook(file_path)
            sheet = workbook.active
            
            if sheet.max_row < 2:
                raise ValueError("Le fichier Excel ne contient pas de données")
            
            # Vérifier les en-têtes
            headers = [cell.value for cell in sheet[1]]
            expected_headers = ["Nom", "Prénom", "Classe", "Année scolaire", "Montant total"]
            
            for expected in expected_headers:
                if expected not in headers:
                    raise ValueError(f"En-tête manquant: {expected}. En-têtes attendus: {', '.join(expected_headers)}")
            
            # Index des colonnes
            col_indices = {
                "nom": headers.index("Nom"),
                "prenom": headers.index("Prénom"),
                "classe": headers.index("Classe"),
                "annee_scolaire": headers.index("Année scolaire"),
                "montant_total": headers.index("Montant total")
            }
            
            imported_count = 0
            error_count = 0
            errors = []
            
            # Parcourir les lignes (à partir de la ligne 2)
            for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                try:
                    nom = str(row[col_indices["nom"]].value).strip() if row[col_indices["nom"]].value else ""
                    prenom = str(row[col_indices["prenom"]].value).strip() if row[col_indices["prenom"]].value else ""
                    classe = str(row[col_indices["classe"]].value).strip() if row[col_indices["classe"]].value else ""
                    annee_scolaire = str(row[col_indices["annee_scolaire"]].value).strip() if row[col_indices["annee_scolaire"]].value else ""
                    montant_total = row[col_indices["montant_total"]].value
                    
                    # Validation
                    if not nom or not prenom or not classe or not annee_scolaire:
                        errors.append(f"Ligne {row_idx}: Champs obligatoires manquants")
                        error_count += 1
                        continue
                    
                    if montant_total is None:
                        errors.append(f"Ligne {row_idx}: Montant total manquant")
                        error_count += 1
                        continue
                    
                    try:
                        montant_total = float(montant_total)
                    except (ValueError, TypeError):
                        errors.append(f"Ligne {row_idx}: Montant total invalide: {montant_total}")
                        error_count += 1
                        continue
                    
                    if montant_total <= 0:
                        errors.append(f"Ligne {row_idx}: Le montant total doit être positif")
                        error_count += 1
                        continue
                    
                    # Créer l'élève
                    self.student_repository.create(
                        nom=nom,
                        prenom=prenom,
                        classe=classe,
                        annee_scolaire=annee_scolaire,
                        montant_total=montant_total
                    )
                    imported_count += 1
                    logger.info(f"Élève importé: {nom} {prenom} ({classe})")
                    
                except Exception as e:
                    error_msg = f"Ligne {row_idx}: {str(e)}"
                    errors.append(error_msg)
                    error_count += 1
                    logger.error(error_msg)
            
            workbook.close()
            return imported_count, error_count, errors
            
        except FileNotFoundError:
            raise FileNotFoundError(f"Fichier non trouvé: {file_path}")
        except Exception as e:
            logger.exception(f"Erreur lors de l'import Excel: {e}")
            raise ValueError(f"Erreur lors de l'import: {str(e)}")
    
    def get_excel_template(self) -> List[List[str]]:
        """Retourne un modèle de structure pour l'import Excel.
        
        Returns:
            Liste de listes représentant les lignes du modèle
        """
        return [
            ["Nom", "Prénom", "Classe", "Année scolaire", "Montant total"],
            ["Dupont", "Jean", "6ème A", "2024-2025", "500000"],
            ["Martin", "Marie", "5ème B", "2024-2025", "450000"],
            ["Bernard", "Pierre", "4ème A", "2024-2025", "600000"]
        ]
